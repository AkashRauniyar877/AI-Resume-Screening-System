"""Test suite for ai_engine package.

Validates PDF parsing, DOCX parsing, preprocessing, information extraction,
embedding generation, matching, hybrid scoring, candidate ranking, and pipeline execution.
"""

import os
import unittest
import numpy as np

from ai_engine.parser import extract_text, extract_text_from_pdf
from ai_engine.preprocessing import clean_raw_text, preprocess_text
from ai_engine.extraction import (
    extract_email,
    extract_phone,
    extract_name,
    extract_skills,
    extract_education,
    extract_experience,
)
from ai_engine.matching import (
    calculate_skill_match,
    compute_tfidf_similarity,
    compute_semantic_similarity,
)
from ai_engine.scoreing import (
    calculate_experience_score,
    min_max_normalize,
    calculate_hybrid_score,
    compute_batch_hybrid_scores,
)
from ai_engine.ranking import rank_candidates
from ai_engine.models import MatchingResult, ResumeInformation
from ai_engine.pipeline import (
    analyze_job_description,
    process_resume,
    process_and_rank_candidates,
    ResumeScreeningPipeline,
)

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))


class TestAIEngine(unittest.TestCase):

    def setUp(self):
        self.sample_pdf = os.path.join(
            PROJECT_ROOT, "data", "resumes", "data", "ACCOUNTANT", "10554236.pdf"
        )
        self.sample_jd = """
        We are looking for a Python Developer with 3+ years of experience
        in backend development.

        The candidate should have strong knowledge of Python, FastAPI,
        REST API, MongoDB, SQL, React and JavaScript.

        Experience with Git, Docker and AWS is preferred.

        The candidate should have a Bachelor's degree in Computer Science,
        Information Technology or a related field.
        """
        self.sample_resume_text = """
        John Doe
        john.doe@example.com
        +1 555-123-4567

        Professional Summary:
        Experienced software engineer with 5 years of experience in Python,
        FastAPI, REST API, SQL, React, Git, and Docker.

        Education:
        Bachelor of Science in Computer Science
        """

    def test_01_pdf_parsing(self):
        """Test PDF parsing on existing sample resume 10554236.pdf."""
        self.assertTrue(
            os.path.exists(self.sample_pdf),
            f"Sample PDF not found at {self.sample_pdf}",
        )
        text = extract_text_from_pdf(self.sample_pdf)
        self.assertIsInstance(text, str)
        self.assertGreater(len(text), 100)
        self.assertIn("ACCOUNTANT", text.upper())

    def test_01b_docx_parsing(self):
        """Test DOCX parsing functionality."""
        from docx import Document
        from ai_engine.parser import extract_text_from_docx

        test_docx_path = os.path.join(PROJECT_ROOT, "scratch_test_resume.docx")
        doc = Document()
        doc.add_paragraph("Alice Smith")
        doc.add_paragraph("alice@example.com")
        doc.add_paragraph("Experienced Python Developer with 4 years experience in Django and SQL.")
        doc.save(test_docx_path)

        try:
            text = extract_text_from_docx(test_docx_path)
            self.assertIn("Alice Smith", text)
            self.assertIn("Django", text)
        finally:
            if os.path.exists(test_docx_path):
                os.remove(test_docx_path)

    def test_02_text_extractor(self):
        """Test text_extractor interface."""
        text = extract_text(self.sample_pdf)
        self.assertIsInstance(text, str)
        self.assertGreater(len(text), 100)

    def test_03_preprocessing(self):
        """Test text cleaning and NLTK stop-word filtering."""
        raw = "<html><body>Hello World! This is a TEST resume.</body></html>"
        cleaned = clean_raw_text(raw)
        self.assertEqual(cleaned, "hello world! this is a test resume.")

        processed = preprocess_text(self.sample_resume_text)
        self.assertIsInstance(processed, str)
        self.assertNotIn("with", processed.split())
        self.assertIn("python", processed.split())

    def test_04_email_extraction(self):
        """Test email extraction matching regex."""
        email = extract_email(self.sample_resume_text)
        self.assertEqual(email, "john.doe@example.com")

    def test_05_phone_extraction(self):
        """Test phone extraction matching regex."""
        phone = extract_phone(self.sample_resume_text)
        self.assertIsNotNone(phone)
        self.assertIn("555-123-4567", phone)

    def test_06_name_extraction(self):
        """Test conservative name extraction."""
        name = extract_name(self.sample_resume_text)
        self.assertEqual(name, "John Doe")

    def test_07_skill_extraction(self):
        """Test skill extraction with overlap pruning."""
        skills = extract_skills(self.sample_resume_text)
        self.assertIn("python", skills)
        self.assertIn("fastapi", skills)
        self.assertIn("rest api", skills)
        # Verify shorter 'rest' or 'api' sub-matches are pruned when 'rest api' is present
        self.assertNotIn("rest", skills)

    def test_08_education_extraction(self):
        """Test degree regex pattern extraction and normalization."""
        edu = extract_education(self.sample_resume_text)
        self.assertIn("Bachelor of Science", edu)

    def test_09_experience_extraction(self):
        """Test explicit experience regex extraction."""
        exp = extract_experience(self.sample_resume_text)
        self.assertEqual(exp, 5.0)

        no_exp = extract_experience("Candidate has software engineering skills.")
        self.assertEqual(no_exp, 0.0)

    def test_10_skill_matching(self):
        """Test candidate skills vs required skills intersection/difference/coverage."""
        cand_skills = ["python", "fastapi", "sql"]
        req_skills = ["python", "fastapi", "react", "aws"]

        matched, missing, coverage = calculate_skill_match(cand_skills, req_skills)
        self.assertCountEqual(matched, ["python", "fastapi"])
        self.assertCountEqual(missing, ["react", "aws"])
        self.assertEqual(coverage, 50.0)

    def test_11_tfidf_matching(self):
        """Test TF-IDF cosine similarity calculations."""
        processed_resumes = [
            "python fastapi rest api sql react git docker",
            "accounting financial ledger excel tax audit",
        ]
        scores = compute_tfidf_similarity(
            "python fastapi developer sql react", processed_resumes
        )
        self.assertEqual(len(scores), 2)
        self.assertGreater(scores[0], scores[1])

    def test_12_semantic_matching(self):
        """Test Sentence Transformer embedding similarity."""
        resumes = [
            "Experienced Python and FastAPI backend developer creating web applications.",
            "Senior Accountant with expertise in balance sheets and tax filings.",
        ]
        scores = compute_semantic_similarity(self.sample_jd, resumes)
        self.assertEqual(len(scores), 2)
        self.assertGreater(scores[0], scores[1])

    def test_13_hybrid_scoring(self):
        """Test min-max normalization and hybrid formula calculation."""
        raw_scores = np.array([20.0, 60.0, 100.0])
        norm = min_max_normalize(raw_scores)
        self.assertAlmostEqual(norm[0], 0.0)
        self.assertAlmostEqual(norm[1], 50.0)
        self.assertAlmostEqual(norm[2], 100.0)

        exp_score = calculate_experience_score(3.0, 3.0)
        self.assertEqual(exp_score, 100.0)

        # Formula: 0.35 * skill + 0.30 * semantic + 0.20 * tfidf + 0.15 * exp
        hybrid = calculate_hybrid_score(
            skill_coverage=100.0,
            semantic_normalized=100.0,
            tfidf_normalized=100.0,
            experience_score=100.0,
        )
        self.assertEqual(hybrid, 100.0)

    def test_14_candidate_ranking(self):
        """Test ranking logic sorting candidates by hybrid score descending."""
        info1 = ResumeInformation(id=101, name="Alice")
        match1 = MatchingResult(candidate_id=101, hybrid_score=85.5)

        info2 = ResumeInformation(id=102, name="Bob")
        match2 = MatchingResult(candidate_id=102, hybrid_score=92.0)

        ranked = rank_candidates([(info1, match1), (info2, match2)])
        self.assertEqual(len(ranked), 2)
        self.assertEqual(ranked[0].candidate_id, 102)
        self.assertEqual(ranked[0].rank, 1)
        self.assertEqual(ranked[1].candidate_id, 101)
        self.assertEqual(ranked[1].rank, 2)

    def test_15_pipeline(self):
        """Test complete pipeline orchestration."""
        jd_analysis = analyze_job_description(self.sample_jd)
        self.assertIn("python", jd_analysis.required_skills)
        self.assertEqual(jd_analysis.required_experience, 3.0)

        resumes = [
            {
                "id": "cand_01",
                "text": "Senior Python Backend Engineer with 5 years experience in Python, FastAPI, REST API, MongoDB, SQL, React, Git, Docker, AWS. Bachelor of Science in CS.",
            },
            {
                "id": "cand_02",
                "text": "Junior Accountant with 1 year experience in Excel, QuickBooks, financial reporting. Bachelor of Commerce.",
            },
        ]

        pipeline = ResumeScreeningPipeline()
        ranked = pipeline.process_and_rank(resumes, self.sample_jd)

        self.assertEqual(len(ranked), 2)
        self.assertEqual(ranked[0].candidate_id, "cand_01")
        self.assertEqual(ranked[0].rank, 1)
        self.assertGreater(ranked[0].hybrid_score, ranked[1].hybrid_score)


if __name__ == "__main__":
    unittest.main()
